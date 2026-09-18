"""Generate SkillSwap AI hackathon presentation (PPTX)."""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt, Emu

ROOT = Path(__file__).resolve().parents[1]
SHOTS = ROOT / "docs" / "screenshots"
OUT = ROOT / "docs" / "SkillSwap_AI_Presentation.pptx"

# Brand
BG = RGBColor(0x0B, 0x0B, 0x0F)
CARD = RGBColor(0x16, 0x16, 0x1D)
ACCENT = RGBColor(0x4F, 0x7F, 0xFF)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
MUTED = RGBColor(0xA8, 0xB0, 0xC4)
SOFT = RGBColor(0x7A, 0xA0, 0xFF)

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)


def set_run(run, size=18, bold=False, color=WHITE, font="Calibri"):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font


def add_bg(slide):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, SLIDE_H)
    shape.fill.solid()
    shape.fill.fore_color.rgb = BG
    shape.line.fill.background()
    # accent bar top
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, Inches(0.08))
    bar.fill.solid()
    bar.fill.fore_color.rgb = ACCENT
    bar.line.fill.background()


def add_text_box(slide, left, top, width, height, text, size=18, bold=False, color=WHITE, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    set_run(run, size=size, bold=bold, color=color)
    return box


def add_bullets(slide, left, top, width, height, items, size=18, color=MUTED):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.level = 0
        p.space_after = Pt(10)
        run = p.add_run()
        run.text = f"•  {item}"
        set_run(run, size=size, color=color)
    return box


def add_card(slide, left, top, width, height):
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    card.fill.solid()
    card.fill.fore_color.rgb = CARD
    card.line.color.rgb = RGBColor(0x2E, 0x2E, 0x3A)
    card.adjustments[0] = 0.1
    return card


def section_title(slide, title, subtitle=None):
    add_text_box(slide, Inches(0.7), Inches(0.35), Inches(12), Inches(0.6), title, size=32, bold=True, color=WHITE)
    if subtitle:
        add_text_box(slide, Inches(0.7), Inches(0.95), Inches(12), Inches(0.4), subtitle, size=16, color=SOFT)


def blank_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    add_bg(slide)
    return slide


def add_image_safe(slide, path, left, top, width=None, height=None):
    path = Path(path)
    if not path.exists():
        return None
    kwargs = {"left": left, "top": top}
    if width:
        kwargs["width"] = width
    if height:
        kwargs["height"] = height
    return slide.shapes.add_picture(str(path), **kwargs)


def build():
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    # 1. Title
    s = blank_slide(prs)
    add_text_box(s, Inches(0.7), Inches(1.8), Inches(12), Inches(0.4), "SMART INDIA HACKATHON  ·  TRACK 2", size=14, bold=True, color=ACCENT)
    add_text_box(s, Inches(0.7), Inches(2.4), Inches(12), Inches(1), "SkillSwap AI", size=54, bold=True, color=WHITE)
    add_text_box(
        s,
        Inches(0.7),
        Inches(3.5),
        Inches(11),
        Inches(1),
        "AI-powered marketplace where students & young creators\nshowcase proof of skill, get matched intelligently, and deliver securely.",
        size=20,
        color=MUTED,
    )
    add_text_box(s, Inches(0.7), Inches(5.4), Inches(12), Inches(0.4), "Not a Fiverr clone  ·  Trust-weighted matching  ·  Milestone-safe workflows", size=16, color=SOFT)
    add_text_box(s, Inches(0.7), Inches(6.5), Inches(12), Inches(0.4), "Live: skillswap-ai-ecru.vercel.app", size=14, color=MUTED)

    # 2. Problem
    s = blank_slide(prs)
    section_title(s, "Problem", "Why traditional freelancing fails students")
    problems = [
        ("Experience gatekeeping", "Marketplaces rank by years of work — students with strong projects get buried."),
        ("Race to the bottom", "Lowest-price ranking pushes underbidding, weak quality, and distrust."),
        ("Weak portfolios", "Raw college projects don't look client-ready without heavy rewrite effort."),
        ("Unsafe delivery", "Flat pay-after-delivery models expose young creators to scope creep and unpaid work."),
    ]
    for i, (title, body) in enumerate(problems):
        col = i % 2
        row = i // 2
        left = Inches(0.7 + col * 6.2)
        top = Inches(1.7 + row * 2.4)
        add_card(s, left, top, Inches(5.9), Inches(2.1))
        add_text_box(s, left + Inches(0.35), top + Inches(0.35), Inches(5.2), Inches(0.5), title, size=20, bold=True, color=SOFT)
        add_text_box(s, left + Inches(0.35), top + Inches(0.95), Inches(5.2), Inches(0.9), body, size=15, color=MUTED)

    # 3. Solution
    s = blank_slide(prs)
    section_title(s, "Solution", "SkillSwap AI — proof-of-skill marketplace")
    solutions = [
        "AI Portfolio Builder turns uploads + notes into client-ready cards",
        "Trust-weighted Talent Match (skills, portfolio, trust — never cheapest first)",
        "Gig marketplace for creators + browse/save for clients",
        "Milestone booking workflow: Pending → Accepted → In Progress → Submitted → Completed",
        "Verified reviews only after completed projects",
        "Realtime chat, notifications, analytics for both roles",
    ]
    add_card(s, Inches(0.7), Inches(1.7), Inches(12), Inches(5))
    add_bullets(s, Inches(1.1), Inches(2.0), Inches(11), Inches(4.4), solutions, size=20, color=WHITE)

    # 4. User Journey
    s = blank_slide(prs)
    section_title(s, "User Journey", "Creator and Client flows")
    # Creator path
    add_card(s, Inches(0.7), Inches(1.7), Inches(5.9), Inches(5))
    add_text_box(s, Inches(1.0), Inches(1.95), Inches(5.3), Inches(0.4), "Creator", size=22, bold=True, color=ACCENT)
    add_bullets(
        s,
        Inches(1.0),
        Inches(2.5),
        Inches(5.3),
        Inches(3.8),
        [
            "Register as Creator",
            "Build AI portfolio cards",
            "Publish gigs",
            "Accept bookings + milestones",
            "Submit delivery → earn trust",
            "Track earnings & analytics",
        ],
        size=16,
        color=MUTED,
    )
    # Client path
    add_card(s, Inches(6.9), Inches(1.7), Inches(5.9), Inches(5))
    add_text_box(s, Inches(7.2), Inches(1.95), Inches(5.3), Inches(0.4), "Client", size=22, bold=True, color=ACCENT)
    add_bullets(
        s,
        Inches(7.2),
        Inches(2.5),
        Inches(5.3),
        Inches(3.8),
        [
            "Register as Client",
            "Post a job brief",
            "Get AI talent matches + reasons",
            "Browse / save creators",
            "Book with milestones",
            "Complete → leave verified review",
        ],
        size=16,
        color=MUTED,
    )

    # 5. Architecture
    s = blank_slide(prs)
    section_title(s, "Architecture", "Production deployment topology")
    boxes = [
        (0.7, "Vercel\nFrontend\nReact + Vite", ACCENT),
        (3.7, "Railway\nBackend\nFastAPI + JWT", SOFT),
        (6.7, "Supabase\nPostgreSQL\nSchema + Seed", MUTED),
        (9.7, "Cloudinary\nMedia\nUploads", WHITE),
    ]
    for left, label, color in boxes:
        add_card(s, Inches(left), Inches(2.3), Inches(2.7), Inches(2.8))
        add_text_box(s, Inches(left + 0.15), Inches(2.9), Inches(2.4), Inches(1.8), label, size=16, bold=True, color=color, align=PP_ALIGN.CENTER)
    add_text_box(s, Inches(0.7), Inches(5.5), Inches(12), Inches(1.2),
                 "HTTPS  →  JWT Auth / REST / WebSocket  →  SQLAlchemy + Alembic  →  OpenAI-compatible AI layer",
                 size=16, color=MUTED, align=PP_ALIGN.CENTER)

    # 6. AI Features
    s = blank_slide(prs)
    section_title(s, "AI Features", "Provider-abstracted OpenAI-compatible service")
    feats = [
        ("generate_portfolio()", "Title, description, skills, tools, portfolio card from uploads + notes"),
        ("match_creators()", "Compatibility score + human reasons using skills, tags, portfolio, trust"),
        ("suggest_pricing()", "Category + experience aware gig pricing guidance"),
        ("chat_assistant()", "Brief writing help for creators and clients"),
    ]
    for i, (title, body) in enumerate(feats):
        top = Inches(1.65 + i * 1.25)
        add_card(s, Inches(0.7), top, Inches(12), Inches(1.1))
        add_text_box(s, Inches(1.0), top + Inches(0.2), Inches(11.4), Inches(0.35), title, size=18, bold=True, color=SOFT)
        add_text_box(s, Inches(1.0), top + Inches(0.55), Inches(11.4), Inches(0.4), body, size=15, color=MUTED)

    # 7. UI Screens
    s = blank_slide(prs)
    section_title(s, "UI Screens", "Premium dark SaaS — Apple × Linear × Stripe inspired")
    images = [
        ("landing.png", 0.5, 1.55, 4.0),
        ("creator-dashboard.png", 4.7, 1.55, 4.0),
        ("ai-builder.png", 8.9, 1.55, 4.0),
    ]
    labels = ["Landing", "Creator Dashboard", "AI Portfolio Builder"]
    for (name, left, top, width), label in zip(images, labels):
        pic = add_image_safe(s, SHOTS / name, Inches(left), Inches(top), width=Inches(width))
        add_text_box(s, Inches(left), Inches(5.85), Inches(width), Inches(0.35), label, size=14, bold=True, color=SOFT, align=PP_ALIGN.CENTER)

    # 7b UI Screens continued
    s = blank_slide(prs)
    section_title(s, "UI Screens", "Client experience")
    images2 = [
        ("client-dashboard.png", 0.5, 1.55, 4.0),
        ("post-job.png", 4.7, 1.55, 4.0),
        ("browse-creators.png", 8.9, 1.55, 4.0),
    ]
    labels2 = ["Client Dashboard", "Post Job + Match", "Browse Creators"]
    for (name, left, top, width), label in zip(images2, labels2):
        add_image_safe(s, SHOTS / name, Inches(left), Inches(top), width=Inches(width))
        add_text_box(s, Inches(left), Inches(5.85), Inches(width), Inches(0.35), label, size=14, bold=True, color=SOFT, align=PP_ALIGN.CENTER)

    # 8. Database
    s = blank_slide(prs)
    section_title(s, "Database", "PostgreSQL on Supabase · SQLAlchemy models + Alembic")
    tables = [
        "Users", "CreatorProfiles", "ClientProfiles", "PortfolioProjects",
        "Gigs", "Jobs", "Bookings", "Milestones",
        "Reviews", "Messages", "Conversations", "Notifications",
        "SavedGigs", "SavedJobs", "SavedCreators",
    ]
    add_card(s, Inches(0.7), Inches(1.7), Inches(12), Inches(5))
    # grid of chips
    for i, t in enumerate(tables):
        col = i % 5
        row = i // 5
        left = Inches(1.0 + col * 2.35)
        top = Inches(2.1 + row * 1.3)
        chip = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, Inches(2.15), Inches(0.7))
        chip.fill.solid()
        chip.fill.fore_color.rgb = RGBColor(0x1C, 0x1C, 0x24)
        chip.line.color.rgb = ACCENT
        chip.adjustments[0] = 0.2
        add_text_box(s, left, top + Inches(0.15), Inches(2.15), Inches(0.45), t, size=13, bold=True, color=WHITE, align=PP_ALIGN.CENTER)

    # 9. Tech Stack
    s = blank_slide(prs)
    section_title(s, "Tech Stack", "Modular · scalable · production-ready")
    stacks = [
        ("Frontend", ["React 18", "Vite + TypeScript", "Tailwind CSS", "React Router", "Framer Motion", "React Query", "Axios"]),
        ("Backend", ["Python 3.11", "FastAPI + Uvicorn", "SQLAlchemy", "Alembic", "JWT + Refresh", "Pydantic", "SlowAPI"]),
        ("Platform", ["PostgreSQL (Supabase)", "Cloudinary", "Vercel", "Railway", "WebSockets", "OpenAI-compatible AI"]),
    ]
    for i, (title, items) in enumerate(stacks):
        left = Inches(0.7 + i * 4.15)
        add_card(s, left, Inches(1.7), Inches(3.95), Inches(5))
        add_text_box(s, left + Inches(0.3), Inches(1.95), Inches(3.4), Inches(0.45), title, size=20, bold=True, color=ACCENT)
        add_bullets(s, left + Inches(0.3), Inches(2.55), Inches(3.4), Inches(3.8), items, size=15, color=MUTED)

    # 10. Business Model
    s = blank_slide(prs)
    section_title(s, "Business Model", "Sustainable marketplace for student talent")
    models = [
        ("Marketplace fee", "Small commission on completed milestone bookings"),
        ("Creator Pro", "Boosted visibility, advanced analytics, AI credits"),
        ("Client teams", "Multi-seat hiring workspaces for campuses & startups"),
        ("Verified badges", "Optional identity / institute verification partnerships"),
    ]
    for i, (title, body) in enumerate(models):
        top = Inches(1.7 + i * 1.25)
        add_card(s, Inches(0.7), top, Inches(12), Inches(1.1))
        add_text_box(s, Inches(1.0), top + Inches(0.2), Inches(11.4), Inches(0.35), title, size=18, bold=True, color=SOFT)
        add_text_box(s, Inches(1.0), top + Inches(0.55), Inches(11.4), Inches(0.4), body, size=15, color=MUTED)

    # 11. Future Scope
    s = blank_slide(prs)
    section_title(s, "Future Scope", "Roadmap beyond SIH MVP")
    future = [
        "Payment gateway + escrow for milestone releases",
        "Vector embeddings for deeper portfolio similarity matching",
        "Mobile apps (React Native) for creators on the go",
        "Institute SSO / campus talent hubs",
        "AI contract assistant + dispute mediation support",
        "Multi-language support for Bharat-scale reach",
    ]
    add_card(s, Inches(0.7), Inches(1.7), Inches(12), Inches(5))
    add_bullets(s, Inches(1.1), Inches(2.1), Inches(11), Inches(4.3), future, size=20, color=WHITE)

    # 12. Thank You
    s = blank_slide(prs)
    add_text_box(s, Inches(0.7), Inches(2.2), Inches(12), Inches(1), "Thank You", size=54, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    add_text_box(s, Inches(0.7), Inches(3.4), Inches(12), Inches(0.5), "SkillSwap AI  ·  Smart India Hackathon Track 2", size=20, color=SOFT, align=PP_ALIGN.CENTER)
    add_text_box(
        s,
        Inches(1.5),
        Inches(4.3),
        Inches(10.3),
        Inches(1.5),
        "Demo  https://skillswap-ai-ecru.vercel.app\nGitHub  https://github.com/devanshgupta2610/SkillSwap-AI\nAPI  https://skillswap-ai-api-production.up.railway.app/docs",
        size=16,
        color=MUTED,
        align=PP_ALIGN.CENTER,
    )
    add_text_box(s, Inches(0.7), Inches(6.3), Inches(12), Inches(0.4), "Questions?", size=18, bold=True, color=ACCENT, align=PP_ALIGN.CENTER)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT)
    print(f"Saved: {OUT}")


if __name__ == "__main__":
    build()
